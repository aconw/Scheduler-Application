from __future__ import annotations
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from email.message import EmailMessage
from datetime import datetime
import pandas as pd
import openpyxl
from scheduler_engine import read_sheet, norm, clean

ROOT=Path(__file__).parent
CONFIG_DEFAULT=ROOT/'config'/'Scheduler_Configuration.xlsx'
ASSETS=ROOT/'assets'

def raw_bytes(upload):
    if upload is None: return None
    if isinstance(upload,(bytes,bytearray)): return bytes(upload)
    if hasattr(upload,'getvalue'): return bytes(upload.getvalue())
    upload.seek(0); return upload.read()

def workbook_info(raw: bytes):
    wb=openpyxl.load_workbook(BytesIO(raw),read_only=False,data_only=True)
    return [(ws.title,ws.max_row,ws.max_column) for ws in wb.worksheets]

def _active(df):
    if 'active' not in df.columns: return df
    v=df['active']
    truth=v.fillna(1).astype(str).str.strip().str.lower().isin(['1','1.0','true','yes','y'])
    return df[truth].copy()

def config_frames(source_bytes=None):
    source=BytesIO(source_bytes) if source_bytes else CONFIG_DEFAULT
    sheets=pd.read_excel(source,sheet_name=None,engine='openpyxl')
    needed=['Training Rules','Locations','Equivalencies','Manual Routing']
    missing=[x for x in needed if x not in sheets]
    if missing: raise ValueError(f'Configuration workbook is missing sheets: {missing}')
    rules=_active(sheets['Training Rules']); locations=sheets['Locations'].copy(); eq=_active(sheets['Equivalencies']); routing=_active(sheets['Manual Routing'])
    rules=rules.where(pd.notna(rules),''); locations=locations.where(pd.notna(locations),''); eq=eq.where(pd.notna(eq),''); routing=routing.where(pd.notna(routing),'')
    return rules,locations,eq,routing

def validate_report(label, raw, required, optional=None):
    if raw is None: return {'label':label,'ok':False,'detail':'Not uploaded','rows':0}
    try:
        rows=read_sheet(raw,required_headers=required,optional_headers=optional or [])
        info=workbook_info(raw)
        return {'label':label,'ok':True,'detail':f'{len(raw):,} bytes • {info} • {len(rows):,} data rows','rows':len(rows)}
    except Exception as e:
        return {'label':label,'ok':False,'detail':str(e),'rows':0}

def workday_export(req, contingent=False):
    template=ASSETS/('Enroll_In_Learning_Content_vContingent_Worker_ID.xlsx' if contingent else 'Enroll_In_Learning_Content_vEmployee.xlsx')
    wb=openpyxl.load_workbook(template); ws=wb['Enroll In Learning Content']
    rows=req[(req['Workday_Ready']=='Yes') & req['Selected_Session_WID'].astype(str).ne('')].copy()
    if contingent: rows=rows[rows['Worker_Type'].map(norm).eq('contingent worker')]
    else: rows=rows[~rows['Worker_Type'].map(norm).eq('contingent worker')]
    for i,(_,r) in enumerate(rows.iterrows(),6):
        ws.cell(i,1,f"{r['Event_Key']}-{i-5}"); ws.cell(i,2,str(r['Selected_Session_WID'])); ws.cell(i,3,str(r['Employee_ID'])); ws.cell(i,4,'Y'); ws.cell(i,5,None)
    bio=BytesIO(); wb.save(bio); return bio.getvalue(),len(rows)

def audit_workbook(result):
    """Operational workbook. No application approval step is required."""
    req=result['requirements'].copy(); seat=result['seat_reservations'].copy()
    selected=req[req['Disposition']=='PROPOSED_SCHEDULE'].copy()
    selected_cols=[
        'Event_Key','Worker_Name','Employee_ID','WID','Worker_Type','Event_Type','Anchor_Date','Position_Title','Job_Code',
        'Cost_Center_ID','Cost_Center_Title','Training_Title','Selected_Session_WID','Selected_Reference_ID','Selected_Start',
        'Selected_End','Selected_Location','Distance_Miles','Scheduling_Policy','Prerequisite','Prerequisite_Status','Explanation'
    ]
    selected=selected[[c for c in selected_cols if c in selected.columns]]
    review=req[req['Disposition'].isin(['REVIEW_REQUIRED','MANUAL_SCHEDULING_REQUIRED'])].copy()
    summary=pd.DataFrame([
        {
            **result.get('source_counts',{}),
            'requirements':len(req),
            'proposed_schedule':int((req['Disposition']=='PROPOSED_SCHEDULE').sum()),
            'review_required':int((req['Disposition']=='REVIEW_REQUIRED').sum()),
            'manual_scheduling_required':int((req['Disposition']=='MANUAL_SCHEDULING_REQUIRED').sum()),
            'previously_completed':int(req['Disposition'].isin(['PREVIOUSLY_COMPLETED','EQUIVALENT_COMPLETION']).sum()),
            'already_enrolled':int((req['Disposition']=='ALREADY_ENROLLED').sum()),
        }
    ])
    bio=BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as xw:
        selected.to_excel(xw,index=False,sheet_name='Selected Sessions')
        req.to_excel(xw,index=False,sheet_name='Requirement Audit')
        review.to_excel(xw,index=False,sheet_name='Review Queue')
        seat.to_excel(xw,index=False,sheet_name='Seat Audit')
        summary.to_excel(xw,index=False,sheet_name='Run Summary')
    bio.seek(0); wb=openpyxl.load_workbook(bio)
    from openpyxl.styles import Font,PatternFill,Alignment
    fill=PatternFill('solid',fgColor='1F4E78'); font=Font(color='FFFFFF',bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
        if ws.max_row>=1:
            for c in ws[1]: c.fill=fill; c.font=font; c.alignment=Alignment(wrap_text=True)
            ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            letter=col[0].column_letter
            vals=list(col)[:200]
            maxlen=max([len(str(c.value)) if c.value is not None else 0 for c in vals]+[8])
            ws.column_dimensions[letter].width=min(maxlen+2,45)
    out=BytesIO(); wb.save(out); return out.getvalue()

def eml_bytes(to_email, subject, body):
    m=EmailMessage(); m['To']=to_email or ''; m['Subject']=subject; m['X-Unsent']='1'; m.set_content(body); return m.as_bytes()

def results_zip(result):
    req=result['requirements']; emp,emp_n=workday_export(req,False); cw,cw_n=workday_export(req,True); audit=audit_workbook(result)
    z=BytesIO()
    with ZipFile(z,'w',ZIP_DEFLATED) as zipf:
        zipf.writestr('Scheduling_Results_and_Audit.xlsx',audit)
        zipf.writestr('Workday/Enroll_In_Learning_Content_Employees.xlsx',emp)
        if cw_n: zipf.writestr('Workday/Enroll_In_Learning_Content_Contingent_Workers.xlsx',cw)
        zipf.writestr('README.txt',
            f'Generated {datetime.now():%Y-%m-%d %H:%M}. Employee Workday rows: {emp_n}. Contingent Worker rows: {cw_n}.\n'
            'Scheduling_Results_and_Audit.xlsx contains Selected Sessions, Requirement Audit, Review Queue, Seat Audit, and Run Summary.\n'
            'The files in the Workday folder retain the supplied Workday upload format.\n')
    return z.getvalue(),emp_n,cw_n

def make_email_drafts(req, routing):
    """Generate drafts directly from the current scheduling result; no approval upload is required."""
    route={norm(r['training_title']):clean(r['recipient_email']) for _,r in routing.iterrows() if clean(r.get('training_title'))}
    z=BytesIO(); manager_count=0
    with ZipFile(z,'w',ZIP_DEFLATED) as zipf:
        # One manager draft per staffing event with at least one selected session.
        for ek in sorted(req.loc[req['Disposition']=='PROPOSED_SCHEDULE','Event_Key'].astype(str).unique()):
            rows=req[req['Event_Key'].astype(str)==ek]
            if rows.empty: continue
            r0=rows.iloc[0]
            lines=[f"Employee: {r0['Worker_Name']}",f"Employee ID: {r0['Employee_ID']}",f"Position: {r0['Position_Title']}",f"Hire / Effective Date: {pd.to_datetime(r0['Anchor_Date']).strftime('%m/%d/%Y') if pd.notna(r0['Anchor_Date']) else ''}",'','Scheduled Training','------------------']
            for _,r in rows[rows['Disposition']=='PROPOSED_SCHEDULE'].sort_values('Selected_Start').iterrows():
                st=pd.to_datetime(r['Selected_Start']) if pd.notna(r['Selected_Start']) else None
                en=pd.to_datetime(r['Selected_End']) if pd.notna(r.get('Selected_End')) else None
                time_text=''
                if st is not None:
                    time_text=st.strftime('%m/%d/%Y %I:%M %p')
                    if en is not None: time_text += ' - ' + en.strftime('%I:%M %p')
                lines.append(f"• {r['Training_Title']} — {time_text} — {r['Selected_Location']}")
            done=rows[rows['Disposition'].isin(['PREVIOUSLY_COMPLETED','EQUIVALENT_COMPLETION'])]
            if not done.empty:
                lines+=['','Previously Completed — No New Assignment','----------------------------------------']
                for _,r in done.iterrows(): lines.append(f"• {r['Training_Title']} — {r.get('Completion_Date','')}")
            body='\n'.join(lines); subject=f"Training Schedule - {r0['Worker_Name']}"
            zipf.writestr(f"Manager Emails/{ek}.eml",eml_bytes(clean(r0['Hiring_Manager_AD']),subject,body)); manager_count+=1
        manual=req[req['Disposition']=='MANUAL_SCHEDULING_REQUIRED']
        for idx,r in manual.iterrows():
            recipient=route.get(norm(r['Training_Title']),''); subject=f"Manual scheduling request - {r['Training_Title']} - {r['Worker_Name']}"
            body=f"Please schedule the following employee for {r['Training_Title']}.\n\nEmployee: {r['Worker_Name']}\nEmployee email: {r.get('Work_Email') or r.get('Home_Email')}\nEmployee ID: {r['Employee_ID']}\nHire / Position Effective Date: {r['Anchor_Date']}\nPosition: {r['Position_Title']}\nCost Center: {r['Cost_Center_Title']} ({r['Cost_Center_ID']})\nHiring Manager: {r['Hiring_Manager_AD']}\n"
            zipf.writestr(f"Manual Scheduling/{idx}.eml",eml_bytes(recipient,subject,body))
    return z.getvalue(),manager_count,len(manual)
