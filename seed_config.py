from pathlib import Path
from datetime import datetime
import sqlite3
import openpyxl
from database import connect, DB_PATH

ROOT=Path(__file__).parent
SRC=ROOT.parent
RULES_FILE=SRC/'Training Documentation - Instructor-led - for automation build.xlsx'
LOCATIONS_FILE=SRC/'Extract_Locations_with_Coordinates.xlsx'


def norm(v):
    return '' if v is None else str(v).strip()

def header_map(ws, row=1):
    return {norm(c.value): i+1 for i,c in enumerate(ws[row]) if norm(c.value)}

def seed_rules(conn):
    if conn.execute('SELECT COUNT(*) FROM training_rules').fetchone()[0] > 0:
        return
    wb=openpyxl.load_workbook(RULES_FILE, read_only=True, data_only=True)
    ws=wb['STANDARD_SCHEDULE Documentation']
    hm=header_map(ws,1)
    now=datetime.now().isoformat(timespec='seconds')
    rows=[]
    for values in ws.iter_rows(min_row=2, values_only=True):
        def g(name):
            idx=hm.get(name)
            return values[idx-1] if idx else None
        title=norm(g('Training Requirements'))
        if not title: continue
        rows.append((norm(g('Job Code')),norm(g('Cost Center')),norm(g('Supervisory Organization')),title,
                     norm(g('Prerequisite')),norm(g('Topic')),norm(g('Scheduling Policy')),norm(g('Timing - modifier')),
                     g('Timing - min (days from anchor)'),g('Timing - max (days from anchor)'),1,'seed',now,'seed',now,1))
    conn.executemany('''INSERT INTO training_rules(job_code,cost_center,supervisory_org,training_title,prerequisite,topic,scheduling_policy,timing_modifier,timing_min,timing_max,active,created_by,created_at,modified_by,modified_at,version)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', rows)
    conn.commit()


def seed_locations(conn):
    if conn.execute('SELECT COUNT(*) FROM locations').fetchone()[0] > 0:
        return
    wb=openpyxl.load_workbook(LOCATIONS_FILE, read_only=True, data_only=True)
    ws=wb.active
    hm=None
    header_row=None
    for ri,row in enumerate(ws.iter_rows(values_only=True),1):
        vals=[norm(x) for x in row]
        if 'Location' in vals and 'Latitude' in vals and 'Longitude' in vals:
            hm={v:i for i,v in enumerate(vals) if v}; header_row=ri; break
    if hm is None: raise RuntimeError('Location headers not found')
    rows=[]
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        def g(name):
            idx=hm.get(name); return row[idx] if idx is not None and idx < len(row) else None
        name=norm(g('Location'))
        if not name: continue
        rows.append((name,norm(g('Reference ID')),norm(g('Address Line 1')),norm(g('Address Line 2')),norm(g('City')),norm(g('Country Region')),norm(g('Postal Code')),norm(g('Country')),g('Latitude'),g('Longitude'),norm(g('Coordinate Precision')),norm(g('Coordinate Source')),norm(g('Inactive'))))
    conn.executemany('''INSERT OR REPLACE INTO locations(location,reference_id,address_line_1,address_line_2,city,region,postal_code,country,latitude,longitude,coordinate_precision,coordinate_source,inactive)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)''', rows)
    conn.commit()

if __name__=='__main__':
    conn=connect()
    seed_rules(conn)
    seed_locations(conn)
    print('db',DB_PATH)
    print('training_rules',conn.execute('SELECT COUNT(*) FROM training_rules').fetchone()[0])
    print('locations',conn.execute('SELECT COUNT(*) FROM locations').fetchone()[0])
    print('equivalencies',conn.execute('SELECT COUNT(*) FROM equivalencies').fetchone()[0])
    print('manual_routing',conn.execute('SELECT COUNT(*) FROM manual_routing').fetchone()[0])
