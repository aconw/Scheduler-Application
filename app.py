from __future__ import annotations
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from email.message import EmailMessage
from datetime import datetime
import pandas as pd
import openpyxl
import streamlit as st

from scheduler_engine import run_scheduler, read_sheet, norm, clean

ROOT=Path(__file__).parent
CONFIG_DEFAULT=ROOT/'config'/'Scheduler_Configuration.xlsx'
ASSETS=ROOT/'assets'

def raw_bytes(upload):
    if upload is None: return None
    if isinstance(upload,(bytes,bytearray)): return bytes(upload)
    if hasattr(upload,'getvalue'): return bytes(upload.getvalue())
    upload.seek(0); return upload.read()

def workbook_info(raw: bytes):
    wb=openpyxl.load_workbook(BytesIO(raw),read_only=False,data_only=True)
    return [(ws.title,ws.max_row,ws.max_column) for ws in wb.worksheets]

def _active(df):
    if 'active' not in df.columns: return df
    v=df['active']
    truth=v.fillna(1).astype(str).str.strip().str.lower().isin(['1','1.0','true','yes','y'])
    return df[truth].copy()

def config_frames(source_bytes=None):
    source=BytesIO(source_bytes) if source_bytes else CONFIG_DEFAULT
    sheets=pd.read_excel(source,sheet_name=None,engine='openpyxl')
    needed=['Training Rules','Locations','Equivalencies','Manual Routing']
    missing=[x for x in needed if x not in sheets]
    if missing: raise ValueError(f'Configuration workbook is missing sheets: {missing}')
    rules=_active(sheets['Training Rules']); locations=sheets['Locations'].copy(); eq=_active(sheets['Equivalencies']); routing=_active(sheets['Manual Routing'])
    rules=rules.where(pd.notna(rules),''); locations=locations.where(pd.notna(locations),''); eq=eq.where(pd.notna(eq),''); routing=routing.where(pd.notna(routing),'')
    return rules,locations,eq,routing

def validate_report(label, raw, required, optional=None):
    if raw is None: return {'label':label,'ok':False,'detail':'Not uploaded','rows':0}
    try:
        rows=read_sheet(raw,required_headers=required,optional_headers=optional or [])
        info=workbook_info(raw)
        return {'label':label,'ok':True,'detail':f'{len(raw):,} bytes • {info} • {len(rows):,} data rows','rows':len(rows)}
    except Exception as e:
        return {'label':label,'ok':False,'detail':str(e),'rows':0}

def workday_export(req, contingent=False):
    """Create a Workday file with hard duplicate protection at person level."""
    template=ASSETS/('Enroll_In_Learning_Content_vContingent_Worker_ID.xlsx' if contingent else 'Enroll_In_Learning_Content_vEmployee.xlsx')
    wb=openpyxl.load_workbook(template); ws=wb['Enroll In Learning Content']
    rows=req[(req['Workday_Ready']=='Yes') & req['Selected_Session_WID'].astype(str).ne('')].copy()
    if contingent: rows=rows[rows['Worker_Type'].map(norm).eq('contingent worker')]
    else: rows=rows[~rows['Worker_Type'].map(norm).eq('contingent worker')]
    if not rows.empty:
        if 'Person_Key' not in rows.columns:
            rows['Person_Key']=rows.apply(lambda r: f"WID:{r.get('WID')}" if clean(r.get('WID')) else f"EID:{r.get('Employee_ID')}",axis=1)
        rows['_course_key']=rows['Training_Title'].map(norm)
        rows['_session_key']=rows['Selected_Session_WID'].astype(str).str.strip()
        # Defense in depth: one person/course and one person/session can appear only once.
        rows=rows.sort_values(['Person_Key','Selected_Start','Event_Key'],na_position='last')
        rows=rows.drop_duplicates(subset=['Person_Key','_course_key'],keep='first')
        rows=rows.drop_duplicates(subset=['Person_Key','_session_key'],keep='first')
    for i,(_,r) in enumerate(rows.iterrows(),6):
        learner=clean(r.get('Current_Employee_ID')) or clean(r.get('Employee_ID'))
        ws.cell(i,1,f"{r['Event_Key']}-{i-5}"); ws.cell(i,2,str(r['Selected_Session_WID'])); ws.cell(i,3,str(learner)); ws.cell(i,4,'Y'); ws.cell(i,5,None)
    bio=BytesIO(); wb.save(bio); return bio.getvalue(),len(rows)

def audit_workbook(result):
    """Operational workbook. No application approval step is required."""
    req=result['requirements'].copy(); seat=result['seat_reservations'].copy()
    duplicate_events=result.get('duplicate_events',pd.DataFrame()).copy()
    selected=req[req['Disposition']=='PROPOSED_SCHEDULE'].copy()
    selected_cols=[
        'Event_Key','Worker_Name','Employee_ID','WID','Worker_Type','Event_Type','Anchor_Date','Position_Title','Job_Code',
        'Cost_Center_ID','Cost_Center_Title','Training_Title','Disposition','Selected_Session_WID','Selected_Reference_ID','Selected_Start',
        'Selected_End','Selected_Location','Distance_Miles','Scheduling_Policy','Prerequisite','Prerequisite_Status','Explanation'
    ]
    selected=selected[[c for c in selected_cols if c in selected.columns]]
    review=req[req['Disposition'].isin(['REVIEW_REQUIRED','MANUAL_SCHEDULING_REQUIRED'])].copy()
    summary=pd.DataFrame([
        {
            **result.get('source_counts',{}),
            'requirements':len(req),
            'proposed_schedule':int((req['Disposition']=='PROPOSED_SCHEDULE').sum()),
            'review_required':int((req['Disposition']=='REVIEW_REQUIRED').sum()),
            'manual_scheduling_required':int((req['Disposition']=='MANUAL_SCHEDULING_REQUIRED').sum()),
            'previously_completed':int(req['Disposition'].isin(['PREVIOUSLY_COMPLETED','EQUIVALENT_COMPLETION']).sum()),
            'already_enrolled':int((req['Disposition']=='ALREADY_ENROLLED').sum()),
            'satisfied_by_same_run_assignment':int((req['Disposition']=='SATISFIED_BY_SAME_RUN_ASSIGNMENT').sum()),
            'duplicate_source_events':int(len(duplicate_events)),
        }
    ])
    bio=BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as xw:
        selected.to_excel(xw,index=False,sheet_name='Selected Sessions')
        req.to_excel(xw,index=False,sheet_name='Requirement Audit')
        review.to_excel(xw,index=False,sheet_name='Review Queue')
        seat.to_excel(xw,index=False,sheet_name='Seat Audit')
        duplicate_events.to_excel(xw,index=False,sheet_name='Duplicate Source Events')
        summary.to_excel(xw,index=False,sheet_name='Run Summary')
    bio.seek(0); wb=openpyxl.load_workbook(bio)
    from openpyxl.styles import Font,PatternFill,Alignment
    fill=PatternFill('solid',fgColor='1F4E78'); font=Font(color='FFFFFF',bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
        if ws.max_row>=1:
            for c in ws[1]: c.fill=fill; c.font=font; c.alignment=Alignment(wrap_text=True)
            ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            letter=col[0].column_letter
            vals=list(col)[:200]
            maxlen=max([len(str(c.value)) if c.value is not None else 0 for c in vals]+[8])
            ws.column_dimensions[letter].width=min(maxlen+2,45)
    out=BytesIO(); wb.save(out); return out.getvalue()

def eml_bytes(to_email, subject, body):
    m=EmailMessage(); m['To']=to_email or ''; m['Subject']=subject; m['X-Unsent']='1'; m.set_content(body); return m.as_bytes()

def results_zip(result):
    req=result['requirements']; emp,emp_n=workday_export(req,False); cw,cw_n=workday_export(req,True); audit=audit_workbook(result)
    z=BytesIO()
    with ZipFile(z,'w',ZIP_DEFLATED) as zipf:
        zipf.writestr('Scheduling_Results_and_Audit.xlsx',audit)
        zipf.writestr('Workday/Enroll_In_Learning_Content_Employees.xlsx',emp)
        if cw_n: zipf.writestr('Workday/Enroll_In_Learning_Content_Contingent_Workers.xlsx',cw)
        zipf.writestr('README.txt',
            f'Generated {datetime.now():%Y-%m-%d %H:%M}. Employee Workday rows: {emp_n}. Contingent Worker rows: {cw_n}.\n'
            'Scheduling_Results_and_Audit.xlsx contains Selected Sessions, Requirement Audit, Review Queue, Seat Audit, Duplicate Source Events, and Run Summary.\n'
            'The files in the Workday folder retain the supplied Workday upload format.\n')
    return z.getvalue(),emp_n,cw_n

def make_email_drafts(req, routing):
    """Generate drafts directly from the current scheduling result; no approval upload is required."""
    route={norm(r['training_title']):clean(r['recipient_email']) for _,r in routing.iterrows() if clean(r.get('training_title'))}
    z=BytesIO(); manager_count=0
    with ZipFile(z,'w',ZIP_DEFLATED) as zipf:
        # One manager draft per staffing event with at least one selected session.
        for ek in sorted(req.loc[req['Disposition']=='PROPOSED_SCHEDULE','Event_Key'].astype(str).unique()):
            rows=req[req['Event_Key'].astype(str)==ek]
            if rows.empty: continue
            r0=rows.iloc[0]
            lines=[f"Employee: {r0['Worker_Name']}",f"Employee ID: {r0['Employee_ID']}",f"Position: {r0['Position_Title']}",f"Hire / Effective Date: {pd.to_datetime(r0['Anchor_Date']).strftime('%m/%d/%Y') if pd.notna(r0['Anchor_Date']) else ''}",'','Scheduled Training','------------------']
            for _,r in rows[rows['Disposition']=='PROPOSED_SCHEDULE'].sort_values('Selected_Start').iterrows():
                st=pd.to_datetime(r['Selected_Start']) if pd.notna(r['Selected_Start']) else None
                en=pd.to_datetime(r['Selected_End']) if pd.notna(r.get('Selected_End')) else None
                time_text=''
                if st is not None:
                    time_text=st.strftime('%m/%d/%Y %I:%M %p')
                    if en is not None: time_text += ' - ' + en.strftime('%I:%M %p')
                lines.append(f"• {r['Training_Title']} — {time_text} — {r['Selected_Location']}")
            done=rows[rows['Disposition'].isin(['PREVIOUSLY_COMPLETED','EQUIVALENT_COMPLETION'])]
            if not done.empty:
                lines+=['','Previously Completed — No New Assignment','----------------------------------------']
                for _,r in done.iterrows(): lines.append(f"• {r['Training_Title']} — {r.get('Completion_Date','')}")
            body='\n'.join(lines); subject=f"Training Schedule - {r0['Worker_Name']}"
            zipf.writestr(f"Manager Emails/{ek}.eml",eml_bytes(clean(r0['Hiring_Manager_AD']),subject,body)); manager_count+=1
        manual=req[req['Disposition']=='MANUAL_SCHEDULING_REQUIRED']
        for idx,r in manual.iterrows():
            recipient=route.get(norm(r['Training_Title']),''); subject=f"Manual scheduling request - {r['Training_Title']} - {r['Worker_Name']}"
            body=f"Please schedule the following employee for {r['Training_Title']}.\n\nEmployee: {r['Worker_Name']}\nEmployee email: {r.get('Work_Email') or r.get('Home_Email')}\nEmployee ID: {r['Employee_ID']}\nHire / Position Effective Date: {r['Anchor_Date']}\nPosition: {r['Position_Title']}\nCost Center: {r['Cost_Center_Title']} ({r['Cost_Center_ID']})\nHiring Manager: {r['Hiring_Manager_AD']}\n"
            zipf.writestr(f"Manual Scheduling/{idx}.eml",eml_bytes(recipient,subject,body))
    return z.getvalue(),manager_count,len(manual)

ROOT=Path(__file__).parent
CONFIG_DEFAULT=ROOT/'config'/'Scheduler_Configuration.xlsx'

st.set_page_config(page_title='Class Scheduling Batch Tool', page_icon='📚', layout='wide')
st.title('Class Scheduling Batch Tool')
st.caption('Simplified stateless build v2.5 • Duplicate source-event collapse • Person-level deduplication • Non-overlapping scheduling')

if 'result' not in st.session_state: st.session_state.result=None
if 'config_bytes' not in st.session_state: st.session_state.config_bytes=None

with st.expander('Configuration',expanded=False):
    st.write('The bundled configuration contains the current training rules and location master. Upload an edited configuration workbook only when you want to use changed rules, equivalencies, locations, or manual-message routing.')
    conf=st.file_uploader('Optional Scheduler_Configuration.xlsx',type=['xlsx'],key='config_upload')
    if conf is not None: st.session_state.config_bytes=raw_bytes(conf)
    st.download_button('Download Current Configuration Workbook',CONFIG_DEFAULT.read_bytes(),'Scheduler_Configuration.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    try:
        cr,cl,ce,cm=config_frames(st.session_state.config_bytes); st.success(f'Configuration ready: {len(cr):,} active rules • {len(cl):,} locations • {len(ce):,} equivalencies • {len(cm):,} manual routes')
    except Exception as e: st.error(str(e))

st.subheader('1. Upload Workday Reports')
c1,c2=st.columns(2)
with c1:
    nh=st.file_uploader('New Hire Orientation Report (optional if Job Change uploaded)',type=['xlsx'],key='nh')
    jc=st.file_uploader('New / Additional Job Change Report (optional if New Hire uploaded)',type=['xlsx'],key='jc')
with c2:
    hist=st.file_uploader('Training History / Learning Transcript (required)',type=['xlsx'],key='hist')
    sess=st.file_uploader('Learning Content / Available Sessions (required)',type=['xlsx'],key='sess')
    orient=st.file_uploader('Existing Orientation Schedule (required — used to prevent conflicts with existing classes)',type=['xlsx'],key='orient')

raws={k:raw_bytes(v) for k,v in {'nh':nh,'jc':jc,'hist':hist,'sess':sess,'orient':orient}.items()}

st.subheader('2. Preflight Validation')
checks=[]
if raws['nh'] is not None: checks.append(validate_report('New Hire',raws['nh'],['Employee ID','WID','Hire Date','Candidate Cost Center ID']))
if raws['jc'] is not None: checks.append(validate_report('Job Change',raws['jc'],['Employee ID','WID','Effective Date','Job Code - Proposed']))
checks.append(validate_report('Training History',raws['hist'],['Employee ID','Record Learning Content','Record Completion Status'],['Learning Participant','WID','Registration Status',"Learner's Registration Date",'Record Completion Date']))
checks.append(validate_report('Available Sessions',raws['sess'],['Learning Content Type','Title','Reference ID','Start Date','End Date','Available Seats','WID']))
checks.append(validate_report('Existing Orientation Schedule',raws['orient'],['Employee ID','Enrolled Course Offering','Registration Status','Start Date','End Date']))
for ck in checks:
    (st.success if ck['ok'] else st.error)(f"{ck['label']}: {ck['detail']}")

required_ok=(raws['hist'] is not None and raws['sess'] is not None and raws['orient'] is not None and (raws['nh'] is not None or raws['jc'] is not None) and all(c['ok'] for c in checks if c['label'] in ['Training History','Available Sessions','Existing Orientation Schedule','New Hire','Job Change']))

if st.button('Run Scheduling',type='primary',use_container_width=True,disabled=not required_ok):
    try:
        rules,locations,eq,routing=config_frames(st.session_state.config_bytes)
        with st.spinner('Scheduling and checking employee time conflicts...'):
            st.session_state.result=run_scheduler(raws['nh'],raws['jc'],raws['hist'],raws['sess'],rules,locations,eq,raws['orient'])
        st.success('Scheduling complete. No selected employee sessions overlap one another. If an Existing Orientation Schedule was uploaded, those classes were also treated as blocked time.')
    except Exception as e: st.exception(e)

if st.session_state.result is not None:
    result=st.session_state.result; req=result['requirements']; counts=req['Disposition'].value_counts().to_dict() if not req.empty and 'Disposition' in req else {}
    st.subheader('3. Results')
    dup_count=len(result.get('duplicate_events',pd.DataFrame()))
    if dup_count:
        st.info(f'{dup_count} duplicate source staffing event(s) were collapsed before training requirements were generated. See Duplicate Source Events in the audit workbook.')
    cols=st.columns(8); metrics=[('Requirements',len(req)),('Selected sessions',counts.get('PROPOSED_SCHEDULE',0)),('Same-run satisfied',counts.get('SATISFIED_BY_SAME_RUN_ASSIGNMENT',0)),('Duplicate events collapsed',dup_count),('Review',counts.get('REVIEW_REQUIRED',0)),('Manual',counts.get('MANUAL_SCHEDULING_REQUIRED',0)),('Completed',counts.get('PREVIOUSLY_COMPLETED',0)+counts.get('EQUIVALENT_COMPLETION',0)),('Already enrolled',counts.get('ALREADY_ENROLLED',0))]
    for col,(lab,val) in zip(cols,metrics): col.metric(lab,val)
    display=['Worker_Name','Employee_ID','Event_Type','Training_Title','Disposition','Selected_Start','Selected_End','Selected_Location','Distance_Miles','Conflict_Detail','Satisfied_By_Event_Key','Satisfied_By_Session_WID','Explanation']
    st.dataframe(req[[c for c in display if c in req.columns]],use_container_width=True,height=440)

    pkg,emp_n,cw_n=results_zip(result)
    st.subheader('4. Export')
    st.write('The package contains exact Workday-format enrollment files plus `Scheduling_Results_and_Audit.xlsx`, which includes **Selected Sessions**, **Requirement Audit**, **Review Queue**, **Seat Audit**, **Duplicate Source Events**, and **Run Summary**. There is no approval/denial step.')
    st.download_button(f'Download Scheduling Export Package ({emp_n} employee / {cw_n} contingent Workday rows)',pkg,'Class_Scheduling_Export_Package.zip',mime='application/zip',type='primary')

    st.subheader('5. Email Drafts (Optional)')
    st.write('Email drafts are generated directly from the current scheduling results. They are drafts only; nothing is sent automatically.')
    try:
        _,_,_,routing=config_frames(st.session_state.config_bytes)
        emails,nmgr,nmanual=make_email_drafts(req,routing)
        st.download_button(f'Download Email Drafts ({nmgr} manager / {nmanual} manual)',emails,'Email_Drafts.zip',mime='application/zip')
    except Exception as e:
        st.error(str(e))
