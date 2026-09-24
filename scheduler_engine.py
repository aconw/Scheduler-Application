from __future__ import annotations
from collections import defaultdict, Counter
from datetime import datetime, date
import math, re, unicodedata
from io import BytesIO
import pandas as pd
import openpyxl

ACTIVE_REGISTRATION_STATUSES={'enrolled','registered','in progress','not started'}
MAX_DISTANCE_MILES=130.0


def clean(v):
    if v is None: return ''
    return str(v).strip() if not isinstance(v,(datetime,date)) else v

def norm(v):
    if v is None: return ''
    return re.sub(r'\s+',' ',str(v).strip().lower())

def header_norm(v):
    """Normalize Workday/Excel headers without altering data values."""
    if v is None: return ''
    s=unicodedata.normalize('NFKC',str(v))
    s=s.replace('\u00a0',' ').replace('\u200b',' ').replace('\r',' ').replace('\n',' ')
    s=s.replace('’',"'").replace('‘',"'")
    s=re.sub(r'\s+',' ',s).strip().lower()
    s=re.sub(r'[:*]+$','',s).strip()
    return s

def id_norm(v):
    if v is None: return ''
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def dt(v):
    if isinstance(v,datetime): return v
    if isinstance(v,date): return datetime(v.year,v.month,v.day)
    if isinstance(v,str) and v.strip():
        s=v.strip()
        for fmt in ('%m/%d/%Y','%Y-%m-%d','%m/%d/%y','%m/%d/%Y %I:%M %p','%Y-%m-%d %H:%M:%S'):
            try: return datetime.strptime(s,fmt)
            except Exception: pass
    return None

def _open_wb(source):
    """Open an uploaded workbook from a stable byte copy.

    Streamlit UploadedFile objects are file-like and their cursor can be moved by
    framework/widget operations. Using getvalue() (or a fresh read after seek)
    prevents an exhausted stream from appearing as an empty workbook.
    """
    if isinstance(source, (bytes, bytearray)):
        raw=bytes(source)
        if not raw:
            raise ValueError('The uploaded Excel file is empty (0 bytes). Please upload the Workday .xlsx export again.')
        return openpyxl.load_workbook(BytesIO(raw), read_only=False, data_only=True)
    if hasattr(source, 'getvalue'):
        raw=source.getvalue()
        if not raw:
            raise ValueError('The uploaded Excel file is empty (0 bytes). Please upload the Workday .xlsx export again.')
        return openpyxl.load_workbook(BytesIO(raw), read_only=False, data_only=True)
    if hasattr(source, 'read'):
        try:
            source.seek(0)
        except Exception:
            pass
        raw=source.read()
        if not raw:
            raise ValueError('The uploaded Excel file could not be read because its upload stream is empty. Please remove and re-upload the file.')
        return openpyxl.load_workbook(BytesIO(raw), read_only=False, data_only=True)
    return openpyxl.load_workbook(source, read_only=False, data_only=True)

# Canonical headers can be recognized under these Workday/export variants.
HEADER_ALIASES={
    'learning participant': {'learning participant','learner','learner name','learning participant name'},
    'employee id': {'employee id','employee_id','employee id number','worker id','employee number'},
    'wid': {'wid','worker wid','employee wid'},
    'record learning content': {'record learning content','learning content','learning content title','record learning content title'},
    'record completion status': {'record completion status','completion status','learning completion status'},
    'record completion date': {'record completion date','completion date','completed date','completed on'},
    'registration status': {'registration status','learner registration status'},
    "learner's registration date": {"learner's registration date",'learners registration date','registration date'},
    'employee id': {'employee id','employee_id','employee id number','worker id','employee number'},
    'enrolled course offering': {'enrolled course offering','course offering','learning enrollment'},
    'start date': {'start date','session start date','offering start date'},
    'learning content type': {'learning content type','content type'},
    'title': {'title','learning content title','course title'},
    'reference id': {'reference id','reference_id'},
    'available seats': {'available seats','seats available','number of available seats'},
    'hire date': {'hire date','most recent hire date'},
    'candidate cost center id': {'candidate cost center id','candidate cost center','cost center id'},
    'effective date': {'effective date','position effective date','staffing event effective date'},
    'job code - proposed': {'job code - proposed','proposed job code','job code proposed'},
}

def _header_matches(value, canonical):
    hv=header_norm(value); ck=header_norm(canonical)
    aliases=HEADER_ALIASES.get(ck,{ck}) | {ck}
    return hv in aliases

def read_sheet(source, sheet=None, header_row=None, required_headers=None, optional_headers=None, scan_rows=150):
    """Read a Workday-style workbook and locate the report headers safely.

    - Uses a fresh byte copy for Streamlit uploads.
    - Searches every worksheet when no sheet is explicitly named.
    - Tolerates Workday title/filter rows and header aliases.
    - Produces actionable diagnostics instead of a generic KeyError.
    """
    wb=_open_wb(source)
    required_headers=list(required_headers or [])
    optional_headers=list(optional_headers or [])

    if sheet is not None:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Expected worksheet '{sheet}' was not found. Available worksheets: {wb.sheetnames}.")
        candidate_sheets=[wb[sheet]]
    else:
        candidate_sheets=list(wb.worksheets)

    workbook_best=(0,None,None,[],[])  # matches, sheet, row, values, missing
    selected=None

    for ws in candidate_sheets:
        if header_row is not None:
            values=[]
            for i,row in enumerate(ws.iter_rows(values_only=True),1):
                if i==header_row:
                    values=list(row); break
            if not values:
                continue
            missing=[req for req in required_headers if not any(_header_matches(v,req) for v in values)]
            matched=[req for req in required_headers if req not in missing]
            if len(matched)>workbook_best[0]: workbook_best=(len(matched),ws.title,header_row,values,missing)
            if not missing:
                selected=(ws,header_row,values); break
        else:
            for i,row in enumerate(ws.iter_rows(values_only=True),1):
                if i>scan_rows: break
                values=list(row)
                matched=[]; missing=[]
                for req in required_headers:
                    if any(_header_matches(v,req) for v in values): matched.append(req)
                    else: missing.append(req)
                if len(matched)>workbook_best[0]: workbook_best=(len(matched),ws.title,i,values,missing)
                if not missing:
                    selected=(ws,i,values); break
            if selected: break

    if selected is None:
        best_count,best_sheet,best_row,best_vals,missing=workbook_best
        detected=[str(v).strip() for v in best_vals if v not in (None,'')]
        # Extra workbook diagnostics help distinguish wrong file vs. empty/blank first tab.
        sheet_stats=[]
        for ws in candidate_sheets:
            nonempty=0
            samples=[]
            for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,10),values_only=True):
                vals=[v for v in row if v not in (None,'')]
                if vals:
                    nonempty+=1
                    if len(samples)<2: samples.append([str(v)[:80] for v in vals[:8]])
            sheet_stats.append(f"{ws.title}: size={ws.max_row}x{ws.max_column}, nonempty rows in first 10={nonempty}, sample={samples}")
        raise ValueError(
            'The uploaded report could not be recognized. '
            f'Expected fields: {required_headers}. '
            f'Best candidate: worksheet={best_sheet or "none"}, row={best_row or "none"}; matched {best_count} of {len(required_headers)}. '
            f'Missing: {missing}. Detected values: {detected[:30]}. '
            f'Workbook worksheets: {wb.sheetnames}. Sheet diagnostics: {sheet_stats}. '
            'Please confirm the correct Workday .xlsx report was uploaded in this field.'
        )

    ws,found_row,header_vals=selected
    raw_headers=list(header_vals)
    canonical_lookup={}
    for canon in required_headers+optional_headers:
        ck=header_norm(canon)
        exact_idx=next((idx for idx,v in enumerate(raw_headers) if header_norm(v)==ck),None)
        if exact_idx is not None:
            canonical_lookup[exact_idx]=canon
            continue
        alias_idx=next((idx for idx,v in enumerate(raw_headers) if _header_matches(v,canon)),None)
        if alias_idx is not None:
            canonical_lookup[alias_idx]=canon

    headers=[]; seen=Counter()
    for c,v in enumerate(raw_headers,1):
        h=canonical_lookup.get(c-1, str(clean(v)) if clean(v)!='' else f'_blank_{c}')
        seen[h]+=1
        if seen[h]>1: h=f'{h}_{seen[h]}'
        headers.append(h)

    out=[]
    for row in ws.iter_rows(min_row=found_row+1,values_only=True):
        if not any(v not in (None,'') for v in row): continue
        vals=list(row)+[None]*(len(headers)-len(row))
        out.append(dict(zip(headers,vals[:len(headers)])))
    return out

def haversine_miles(lat1,lon1,lat2,lon2):
    R=3958.7613
    p1=math.radians(lat1); p2=math.radians(lat2)
    dphi=math.radians(lat2-lat1); dl=math.radians(lon2-lon1)
    a=math.sin(dphi/2)**2+math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return 2*R*math.asin(math.sqrt(a))

def contingent_action(worker_type, traveler):
    if norm(worker_type)!='contingent worker': return 'STANDARD SCHEDULE'
    tr=norm(traveler)
    if tr=='trav i': return 'STANDARD SCHEDULE'
    if tr=='trav d': return 'ESCALATE'
    return 'NO SCHEDULE'

def timing_ok(rec, start_dt):
    anchor=rec.get('Anchor_Date')
    if not anchor or not start_dt: return True
    pol=norm(rec.get('Scheduling_Policy'))
    if pol!='target_range': return start_dt >= anchor
    mn=rec.get('Timing_Min'); mx=rec.get('Timing_Max'); mod=norm(rec.get('Timing_Modifier'))
    try: mn=float(mn) if mn not in (None,'') else None
    except Exception: mn=None
    try: mx=float(mx) if mx not in (None,'') else None
    except Exception: mx=None
    delta=(start_dt.date()-anchor.date()).days
    if mod=='between': return (mn is None or delta>=mn) and (mx is None or delta<=mx)
    if mod=='after':
        t=mn if mn is not None else mx; return t is None or delta>=t
    if mod=='before':
        t=mx if mx is not None else mn; return t is None or delta<=t
    return (mn is None or delta>=mn) and (mx is None or delta<=mx)

def intervals_overlap(start_a, end_a, start_b, end_b):
    """Half-open interval overlap test. Back-to-back sessions are allowed."""
    return start_a < end_b and end_a > start_b

def _person_key(wid, employee_id):
    wid=id_norm(wid); employee_id=id_norm(employee_id)
    return f'WID:{wid}' if wid else f'EID:{employee_id}'

def _event_signature(ev):
    """Identity for duplicate source staffing rows.

    Two rows are duplicates only when the same person, event type, anchor date,
    job/position context, organization, and physical location all match. This
    deliberately runs before requirement generation so a duplicated Workday
    source row cannot build a second prerequisite chain.
    """
    anchor=ev.get('Anchor_Date')
    if isinstance(anchor,datetime): anchor=anchor.isoformat()
    elif isinstance(anchor,date): anchor=anchor.isoformat()
    else: anchor=str(anchor or '')
    return (
        ev.get('Person_Key') or _person_key(ev.get('WID'),ev.get('Employee_ID')),
        norm(ev.get('Event_Type')), anchor,
        norm(ev.get('Job_Code')), norm(ev.get('Position_Title')),
        norm(ev.get('Cost_Center_ID')), norm(ev.get('Sup_Org_ID')),
        norm(ev.get('Physical_Location')), norm(ev.get('Worker_Type')),
        norm(ev.get('Traveler_Designation')),
    )

def run_scheduler(new_hire_file, job_change_file, history_file, sessions_file, rules_df, locations_df, equivalencies_df=None, orientation_file=None):
    new_hires=read_sheet(new_hire_file,required_headers=['Employee ID','WID','Hire Date','Candidate Cost Center ID']) if new_hire_file else []
    job_changes=read_sheet(job_change_file,required_headers=['Employee ID','WID','Effective Date','Job Code - Proposed']) if job_change_file else []
    # WID and Learning Participant are intentionally optional. Employee ID is enough to use the history safely.
    history=read_sheet(
        history_file,
        required_headers=['Employee ID','Record Learning Content','Record Completion Status'],
        optional_headers=['Learning Participant','WID','Registration Status',"Learner's Registration Date",'Record Completion Date']
    )
    sessions=read_sheet(sessions_file,required_headers=['Learning Content Type','Title','Reference ID','Start Date','End Date','Available Seats','WID'])
    orientation=read_sheet(orientation_file,required_headers=['Employee ID','Enrolled Course Offering','Registration Status','Start Date','End Date']) if orientation_file else []

    loc_by_name={}
    for _,r in locations_df.iterrows():
        name=clean(r.get('location'))
        if not name: continue
        try: lat=float(r.get('latitude')); lon=float(r.get('longitude'))
        except Exception: lat=lon=None
        loc_by_name[norm(name)]={'name':name,'lat':lat,'lon':lon,'inactive':clean(r.get('inactive'))}
    def dist(a,b):
        x=loc_by_name.get(norm(a)); y=loc_by_name.get(norm(b))
        if not x or not y or x['lat'] is None or y['lat'] is None: return None
        return haversine_miles(x['lat'],x['lon'],y['lat'],y['lon'])

    emp_to_hist_wid={}
    for h in history:
        if id_norm(h.get('Employee ID')) and id_norm(h.get('WID')):
            emp_to_hist_wid[id_norm(h.get('Employee ID'))]=id_norm(h.get('WID'))

    events=[]
    for i,r in enumerate(new_hires,start=1):
        eid=id_norm(r.get('Employee ID')); source_wid=id_norm(r.get('WID'))
        events.append({'Event_Key':f'NH-{i}-{eid}','Event_Type':'NEW_HIRE','Employee_ID':eid,'WID':emp_to_hist_wid.get(eid,source_wid),'Source_WID':source_wid,
            'Worker_Name':clean(r.get('Worker')),'Worker_Type':clean(r.get('Worker Type')),'Traveler_Designation':clean(r.get('Traveler Designation')),
            'Anchor_Date':dt(r.get('Hire Date')),'Job_Code':clean(r.get('Candidate Position ID')),'Position_Title':clean(r.get('New Hire Position Title')),
            'Cost_Center_ID':clean(r.get('Candidate Cost Center ID')),'Cost_Center_Title':clean(r.get('Candidate Department Name')),
            'Sup_Org_ID':clean(r.get('Supervisory Organization ID')),'Physical_Location':clean(r.get('Worker Physical Location')),
            'Hiring_Manager_AD':clean(r.get('Hiring Manager AD (Username)')),'Work_Email':clean(r.get('Email - Primary Work')),'Home_Email':clean(r.get('Email - Primary Home'))})
    for i,r in enumerate(job_changes,start=1):
        eid=id_norm(r.get('Employee ID')); source_wid=id_norm(r.get('WID'))
        events.append({'Event_Key':f'JC-{i}-{eid}','Event_Type':'JOB_CHANGE','Employee_ID':eid,'WID':emp_to_hist_wid.get(eid,source_wid),'Source_WID':source_wid,
            'Worker_Name':clean(r.get('Worker')),'Worker_Type':clean(r.get('Worker Type')),'Traveler_Designation':'',
            'Anchor_Date':dt(r.get('Effective Date')),'Job_Code':clean(r.get('Job Code - Proposed')),'Position_Title':clean(r.get('Job Title - Proposed')),
            'Cost_Center_ID':clean(r.get('Cost Center ID - Requested')),'Cost_Center_Title':clean(r.get('Cost Center - Proposed')),
            'Sup_Org_ID':clean(r.get('Sup Org ID - Requested')),'Physical_Location':clean(r.get('Location - Proposed')),
            'Hiring_Manager_AD':clean(r.get('Hiring Manager AD (Username')) or clean(r.get('Hiring Manager AD (Username)')),
            'Work_Email':clean(r.get('Candidate Email')),'Home_Email':''})
    for ev in events: ev['Contingent_Action']=contingent_action(ev['Worker_Type'],ev['Traveler_Designation'])

    # Durable person identity and current business identifier. WID is the preferred
    # person key. If a person's Employee ID changes, use the ID from their most
    # recent staffing event for Workday export while retaining each event's source ID.
    latest_id_by_person={}
    for ev in events:
        pkey=_person_key(ev.get('WID'),ev.get('Employee_ID'))
        stamp=ev.get('Anchor_Date') or datetime.min
        prior=latest_id_by_person.get(pkey)
        if prior is None or stamp >= prior[0]:
            latest_id_by_person[pkey]=(stamp,id_norm(ev.get('Employee_ID')))
    for ev in events:
        ev['Person_Key']=_person_key(ev.get('WID'),ev.get('Employee_ID'))
        ev['Current_Employee_ID']=latest_id_by_person.get(ev['Person_Key'],(None,id_norm(ev.get('Employee_ID'))))[1] or id_norm(ev.get('Employee_ID'))

    # Collapse exact duplicate staffing rows BEFORE training requirements are
    # generated. Distinct positions/events are preserved; only rows whose
    # person + event + role/org/location signature is identical are removed.
    raw_event_count=len(events)
    canonical_events=[]; duplicate_events=[]; event_by_signature={}
    for ev in events:
        sig=_event_signature(ev)
        prior=event_by_signature.get(sig)
        if prior is None:
            ev['Event_Status']='CANONICAL'
            ev['Canonical_Event_Key']=ev['Event_Key']
            event_by_signature[sig]=ev
            canonical_events.append(ev)
        else:
            dup=dict(ev)
            dup['Event_Status']='DUPLICATE_SOURCE_EVENT'
            dup['Canonical_Event_Key']=prior['Event_Key']
            duplicate_events.append(dup)
    events=canonical_events

    rule_rows=[]
    for _,r in rules_df.iterrows():
        title=clean(r.get('training_title'))
        if not title: continue
        rule_rows.append({'Rule_ID':r.get('id'),'Job_Code':clean(r.get('job_code')),'Cost_Center':clean(r.get('cost_center')),'Sup_Org':clean(r.get('supervisory_org')),
            'Training_Title':title,'Prerequisite':clean(r.get('prerequisite')),'Topic':clean(r.get('topic')),'Scheduling_Policy':clean(r.get('scheduling_policy')),
            'Timing_Modifier':clean(r.get('timing_modifier')),'Timing_Min':r.get('timing_min'),'Timing_Max':r.get('timing_max')})

    # Equivalency relationships are directional unless explicitly marked TWO_WAY.
    # ONE_WAY semantics: equivalent_training_title satisfies required_training_title.
    # TWO_WAY semantics: either configured title satisfies the other.
    eq=defaultdict(set)
    eq_meta={}
    if equivalencies_df is not None and not equivalencies_df.empty:
        for _,r in equivalencies_df.iterrows():
            req_raw=clean(r.get('required_training_title')); equiv_raw=clean(r.get('equivalent_training_title'))
            req=norm(req_raw); equiv=norm(equiv_raw)
            direction=norm(r.get('relationship_direction') or 'ONE_WAY').replace('-','_').replace(' ','_')
            if direction in ('1_way','oneway'): direction='one_way'
            if direction in ('2_way','twoway'): direction='two_way'
            if direction not in ('one_way','two_way'): direction='one_way'
            if req and equiv:
                eq[req].add(equiv)
                eq_meta[(req,equiv)]={
                    'direction':'ONE_WAY' if direction=='one_way' else 'TWO_WAY',
                    'configured_required':req_raw,'configured_equivalent':equiv_raw,
                    'match_direction':'FORWARD'
                }
                if direction=='two_way':
                    eq[equiv].add(req)
                    eq_meta[(equiv,req)]={
                        'direction':'TWO_WAY',
                        'configured_required':req_raw,'configured_equivalent':equiv_raw,
                        'match_direction':'REVERSE'
                    }

    def satisfaction_titles(required_title):
        key=norm(required_title)
        return {key} | set(eq.get(key,set()))

    def equivalency_match(required_title, observed_title):
        rk=norm(required_title); ok=norm(observed_title)
        if not rk or not ok or rk==ok: return None
        return eq_meta.get((rk,ok))

    hist_by_wid=defaultdict(list); hist_by_eid=defaultdict(list)
    for h in history:
        if id_norm(h.get('WID')): hist_by_wid[id_norm(h.get('WID'))].append(h)
        if id_norm(h.get('Employee ID')): hist_by_eid[id_norm(h.get('Employee ID'))].append(h)
    orient_by_eid=defaultdict(list)
    for o in orientation:
        if id_norm(o.get('Employee ID')): orient_by_eid[id_norm(o.get('Employee ID'))].append(o)
    sessions_by_title=defaultdict(list)
    for s in sessions:
        if clean(s.get('Title')): sessions_by_title[norm(s.get('Title'))].append(s)

    # Session-level view of active/upcoming existing Workday enrollments. The source
    # is lesson-level, so duplicate lesson rows for the same person/session are collapsed.
    existing_workday_rows=[]
    existing_seen=set()

    # Employee time constraints. Existing scheduled/enrolled classes are blocked before
    # any new selections are made. New selections are added as the run proceeds, so
    # one employee can never receive two overlapping sessions, even across staffing events.
    eid_to_wid={}
    for ev in events:
        if ev.get('Employee_ID') and ev.get('WID'):
            eid_to_wid[id_norm(ev['Employee_ID'])]=id_norm(ev['WID'])
    for eid,wid in emp_to_hist_wid.items():
        if eid and wid: eid_to_wid[id_norm(eid)]=id_norm(wid)
    busy_by_person=defaultdict(list)
    for o in orientation:
        if norm(o.get('Registration Status')) not in ACTIVE_REGISTRATION_STATUSES:
            continue
        st=dt(o.get('Start Date')); en=dt(o.get('End Date'))
        if not st or not en or en <= st:
            continue
        eid=id_norm(o.get('Employee ID')); pkey=_person_key(eid_to_wid.get(eid,''),eid)
        title=clean(o.get('Enrolled Course Offering'))
        okey=(pkey,norm(title),st.isoformat(),en.isoformat())
        if okey not in existing_seen:
            existing_seen.add(okey)
            existing_workday_rows.append({
                'Person_Key':pkey,'WID':eid_to_wid.get(eid,''),'Employee_ID':eid,
                'Training_Title':title,'Registration_Status':clean(o.get('Registration Status')),
                'Start_Date':st,'End_Date':en,'Location':clean(o.get('Locations') or o.get('Training Room') or ''),
                'Course_Offering':title,
            })
        busy_by_person[pkey].append({
            'start':st,'end':en,'source':'EXISTING_SCHEDULE',
            'title':title,'event_key':'',
        })

    reqs=[]; seen=set()
    for ev in events:
        if norm(ev['Contingent_Action'])=='no schedule': continue
        for rr in rule_rows:
            if norm(rr['Job_Code'])!=norm(ev['Job_Code']) or norm(rr['Cost_Center'])!=norm(ev['Cost_Center_ID']): continue
            if rr['Sup_Org'] and norm(rr['Sup_Org'])!=norm(ev['Sup_Org_ID']): continue
            key=(ev['Event_Key'],norm(rr['Training_Title']))
            if key in seen: continue
            seen.add(key); reqs.append({**ev,**rr})

    person_req_titles=defaultdict(set)
    for r in reqs:
        person_req_titles[_person_key(r.get('WID'),r.get('Employee_ID'))] |= satisfaction_titles(r['Training_Title'])
    for ew in existing_workday_rows:
        ew['Required_For_Current_Role']='Yes' if norm(ew['Training_Title']) in person_req_titles.get(ew['Person_Key'],set()) else 'No'
        ew['Scheduling_Impact']='BLOCKS_OVERLAP_CHECK'

    def skey(s): return id_norm(s.get('WID')) or f"{clean(s.get('Reference ID'))}|{clean(s.get('Start Date'))}|{norm(s.get('Title'))}"
    opening={}; remaining={}
    for s in sessions:
        try: seats=max(0,int(float(s.get('Available Seats') or 0)))
        except Exception: seats=0
        k=skey(s); opening[k]=max(opening.get(k,0),seats); remaining[k]=opening[k]
    reservations=[]

    for rec in reqs:
        rec.update({'Completion_Date':'','Completion_Training':'','Existing_Registration_Date':'','Existing_Session_Start':'','Existing_Enrollment_Training':'','Equivalency_Used':'','Equivalency_Direction':'','Equivalency_Match_Direction':'','Selected_Session_WID':'','Selected_Reference_ID':'','Selected_Start':'','Selected_End':'','Selected_Location':'','Distance_Miles':'','Original_Available_Seats':'','Remaining_Seats_After_Reservation':'','Prerequisite_Status':'','Prerequisite_Scheduled_Start':'','Conflict_Detail':'','Satisfied_By_Event_Key':'','Satisfied_By_Session_WID':'','Disposition':'','Explanation':'','Workday_Ready':'No','Override':'No','Override_Reason':''})
        title_n=norm(rec['Training_Title']); person_hist=hist_by_wid.get(rec['WID'],[]) or hist_by_eid.get(rec['Employee_ID'],[])
        satisfy_titles=satisfaction_titles(rec['Training_Title'])
        completed=[h for h in person_hist if norm(h.get('Record Learning Content')) in satisfy_titles and norm(h.get('Record Completion Status'))=='completed']
        if completed:
            latest=max(completed,key=lambda h:dt(h.get('Record Completion Date')) or datetime.min)
            observed=clean(latest.get('Record Learning Content'))
            em=equivalency_match(rec['Training_Title'],observed)
            rec['Disposition']='PREVIOUSLY_COMPLETED' if norm(observed)==title_n else 'EQUIVALENT_COMPLETION'
            rec['Completion_Date']=latest.get('Record Completion Date'); rec['Completion_Training']=observed
            if em:
                rec['Equivalency_Used']=f"{em['configured_equivalent']} -> {em['configured_required']}" if em['match_direction']=='FORWARD' else f"{em['configured_required']} <-> {em['configured_equivalent']}"
                rec['Equivalency_Direction']=em['direction']; rec['Equivalency_Match_Direction']=em['match_direction']
                rec['Explanation']=f"Historical completion of equivalent training '{observed}' satisfies required training '{rec['Training_Title']}' via {em['direction']} equivalency. No reassignment required."
            else:
                rec['Explanation']='Historical completion satisfies requirement. No reassignment required.'
            continue

        # Existing enrollment suppression uses the same equivalency relationships as completion suppression.
        existing=[h for h in person_hist if norm(h.get('Record Learning Content')) in satisfy_titles and norm(h.get('Registration Status')) in ACTIVE_REGISTRATION_STATUSES]
        orient_existing=[o for o in orient_by_eid.get(rec['Employee_ID'],[]) if norm(o.get('Enrolled Course Offering')) in satisfy_titles and norm(o.get('Registration Status')) in ACTIVE_REGISTRATION_STATUSES]
        if existing or orient_existing:
            starts=[dt(o.get('Start Date')) for o in orient_existing if dt(o.get('Start Date'))]
            if starts: rec['Existing_Session_Start']=min(starts)
            dates=[dt(h.get("Learner's Registration Date")) for h in existing if dt(h.get("Learner's Registration Date"))]
            if dates: rec['Existing_Registration_Date']=max(dates)
            # Prefer a dated orientation enrollment for the audit; otherwise use the most recent transcript registration.
            observed=''
            if orient_existing:
                dated=[o for o in orient_existing if dt(o.get('Start Date'))]
                chosen=min(dated,key=lambda o:dt(o.get('Start Date'))) if dated else orient_existing[0]
                observed=clean(chosen.get('Enrolled Course Offering'))
            elif existing:
                dated=[h for h in existing if dt(h.get("Learner's Registration Date"))]
                chosen=max(dated,key=lambda h:dt(h.get("Learner's Registration Date"))) if dated else existing[0]
                observed=clean(chosen.get('Record Learning Content'))
            rec['Existing_Enrollment_Training']=observed
            em=equivalency_match(rec['Training_Title'],observed)
            rec['Disposition']='ALREADY_ENROLLED'
            if em:
                rec['Equivalency_Used']=f"{em['configured_equivalent']} -> {em['configured_required']}" if em['match_direction']=='FORWARD' else f"{em['configured_required']} <-> {em['configured_equivalent']}"
                rec['Equivalency_Direction']=em['direction']; rec['Equivalency_Match_Direction']=em['match_direction']
                rec['Explanation']=f"Existing active enrollment in equivalent training '{observed}' satisfies required training '{rec['Training_Title']}' via {em['direction']} equivalency. Duplicate enrollment suppressed."
            else:
                rec['Explanation']='Existing active registration found. Duplicate enrollment suppressed.'
            continue
        if norm(rec['Contingent_Action'])=='escalate': rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']='Contingent worker policy requires escalation.'; continue
        pol=norm(rec['Scheduling_Policy'])
        if pol=='flag_manual': rec['Disposition']='MANUAL_SCHEDULING_REQUIRED'; rec['Explanation']='Training rule uses FLAG_MANUAL.'; continue
        if not pol: rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']='Training rule has no scheduling policy.'; continue
        rec['Disposition']='PENDING_SCHEDULING'

    req_event=defaultdict(dict)
    for r in reqs: req_event[r['Event_Key']][norm(r['Training_Title'])]=r
    def depth(rec,memo,active=None):
        k=norm(rec['Training_Title'])
        if k in memo: return memo[k]
        active=set() if active is None else set(active)
        if k in active: return 999
        p=norm(rec.get('Prerequisite'))
        if not p or p not in req_event[rec['Event_Key']]: memo[k]=0; return 0
        active.add(k); memo[k]=1+depth(req_event[rec['Event_Key']][p],memo,active); return memo[k]

    # One enrollment per person + training course per scheduling run.
    # Additional staffing events that require the same course remain in the audit,
    # but they reference the first assignment instead of consuming another seat.
    same_run_assignment={}

    for ev in sorted(events,key=lambda e:(e.get('Anchor_Date') or datetime.max,e['Event_Key'])):
        er=list(req_event.get(ev['Event_Key'],{}).values()); memo={}; er.sort(key=lambda r:(depth(r,memo),norm(r['Training_Title'])))
        for rec in er:
            if rec['Disposition']!='PENDING_SCHEDULING': continue
            prereq=norm(rec.get('Prerequisite')); prereq_bound=None
            if prereq:
                pr=req_event[rec['Event_Key']].get(prereq)
                if pr:
                    rec['Prerequisite_Status']=pr['Disposition']
                    if pr['Disposition'] in ('PREVIOUSLY_COMPLETED','EQUIVALENT_COMPLETION'): rec['Prerequisite_Status']='SATISFIED_BY_COMPLETION'
                    elif pr['Disposition']=='PROPOSED_SCHEDULE' and dt(pr.get('Selected_Start')): prereq_bound=dt(pr['Selected_Start']); rec['Prerequisite_Scheduled_Start']=pr['Selected_Start']
                    elif pr['Disposition']=='ALREADY_ENROLLED' and dt(pr.get('Existing_Session_Start')): prereq_bound=dt(pr['Existing_Session_Start']); rec['Prerequisite_Scheduled_Start']=pr['Existing_Session_Start']
                    else: rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']=f"Prerequisite '{rec.get('Prerequisite')}' is not in a schedulable state/date ({pr['Disposition']})."; continue
                else:
                    person_hist=hist_by_wid.get(rec['WID'],[]) or hist_by_eid.get(rec['Employee_ID'],[])
                    prereq_titles=satisfaction_titles(rec.get('Prerequisite'))
                    ph=[h for h in person_hist if norm(h.get('Record Learning Content')) in prereq_titles and norm(h.get('Record Completion Status'))=='completed']
                    if ph: rec['Prerequisite_Status']='SATISFIED_BY_PRIOR_COMPLETION'
                    else:
                        pe=[o for o in orient_by_eid.get(rec['Employee_ID'],[]) if norm(o.get('Enrolled Course Offering')) in prereq_titles and norm(o.get('Registration Status')) in ACTIVE_REGISTRATION_STATUSES and dt(o.get('Start Date'))]
                        if pe: prereq_bound=min(dt(o.get('Start Date')) for o in pe); rec['Prerequisite_Status']='SATISFIED_BY_EXISTING_ENROLLMENT'; rec['Prerequisite_Scheduled_Start']=prereq_bound
                        else: rec['Disposition']='REVIEW_REQUIRED'; rec['Prerequisite_Status']='PREREQUISITE_NOT_FOUND'; rec['Explanation']=f"Prerequisite '{rec.get('Prerequisite')}' has no prior completion or dated active enrollment (including configured equivalents)."; continue

            # If this person already received this course in this run, do not
            # create a second enrollment. The existing assignment must still satisfy
            # this staffing event's timing and prerequisite constraints; otherwise
            # surface the second event for review rather than double-enroll the person.
            pkey=_person_key(rec.get('WID'),rec.get('Employee_ID'))
            course_key=(pkey,norm(rec['Training_Title']))
            prior_assignment=same_run_assignment.get(course_key)
            if prior_assignment is not None:
                prior_start=dt(prior_assignment.get('Selected_Start'))
                valid_for_event=bool(prior_start and (not rec.get('Anchor_Date') or prior_start >= rec['Anchor_Date']) and timing_ok(rec,prior_start))
                if prereq_bound and prior_start:
                    valid_for_event=valid_for_event and prior_start > prereq_bound
                if valid_for_event:
                    for fld in ('Selected_Session_WID','Selected_Reference_ID','Selected_Start','Selected_End','Selected_Location','Distance_Miles'):
                        rec[fld]=prior_assignment.get(fld,'')
                    rec['Disposition']='SATISFIED_BY_SAME_RUN_ASSIGNMENT'
                    rec['Workday_Ready']='No'
                    rec['Satisfied_By_Event_Key']=prior_assignment.get('Event_Key','')
                    rec['Satisfied_By_Session_WID']=prior_assignment.get('Selected_Session_WID','')
                    rec['Explanation']=f"Same person/course was already assigned once in this run under staffing event {prior_assignment.get('Event_Key','')}; this requirement is satisfied by that single enrollment."
                else:
                    rec['Disposition']='REVIEW_REQUIRED'
                    rec['Workday_Ready']='No'
                    rec['Satisfied_By_Event_Key']=prior_assignment.get('Event_Key','')
                    rec['Satisfied_By_Session_WID']=prior_assignment.get('Selected_Session_WID','')
                    rec['Explanation']="This person already has the same course assigned once in this run, but that shared session does not satisfy this staffing event's timing/prerequisite rule. A second enrollment was not created."
                continue

            all_s=sessions_by_title.get(norm(rec['Training_Title']),[])
            if not all_s: rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']='No session title matches this training requirement.'; continue
            candidates=[]; overlap_rejected=[]; structurally_eligible=[]
            pkey=_person_key(rec.get('WID'),rec.get('Employee_ID'))
            for s in all_s:
                st=dt(s.get('Start Date')); en=dt(s.get('End Date')); k=skey(s)
                if norm(s.get('Availability Status'))!='open' or not st or not en or en <= st: continue
                if rec['Anchor_Date'] and st < rec['Anchor_Date']: continue
                if prereq_bound and st <= prereq_bound: continue
                if not timing_ok(rec,st): continue
                structurally_eligible.append(s)
                if remaining.get(k,0)<=0: continue
                conflicts=[b for b in busy_by_person.get(pkey,[]) if intervals_overlap(st,en,b['start'],b['end'])]
                if conflicts:
                    overlap_rejected.append((s,conflicts))
                    continue
                candidates.append(s)
            if not candidates:
                capacity_available=[s for s in structurally_eligible if remaining.get(skey(s),0)>0]
                exhausted=[s for s in structurally_eligible if opening.get(skey(s),0)>0 and remaining.get(skey(s),0)<=0]
                rec['Disposition']='REVIEW_REQUIRED'
                if capacity_available and overlap_rejected and len(overlap_rejected) >= len(capacity_available):
                    titles=[]
                    for _,conflicts in overlap_rejected[:5]:
                        for b in conflicts:
                            label=f"{b['title']} ({b['start']:%m/%d/%Y %I:%M %p}-{b['end']:%I:%M %p})"
                            if label not in titles: titles.append(label)
                    rec['Conflict_Detail']='; '.join(titles)
                    rec['Explanation']='All otherwise eligible sessions overlap an existing or newly selected class for this employee.'
                elif exhausted:
                    rec['Explanation']='Eligible session capacity was consumed by earlier assignments in this run.'
                else:
                    rec['Explanation']='No open session with remaining seats meets date/policy/prerequisite/non-overlap criteria.'
                continue
            physical=[s for s in candidates if clean(s.get('Locations'))]; virtual=[s for s in candidates if not clean(s.get('Locations'))]
            mapped=[]
            for s in physical:
                d=dist(rec['Physical_Location'],clean(s.get('Locations')))
                if d is not None: mapped.append((d,s))
            within=[(d,s) for d,s in mapped if d<=MAX_DISTANCE_MILES]
            chosen=None; chosen_dist=''; explanation=''
            if within:
                nearest=min(d for d,_ in within)
                nearest_loc=norm(min((s for d,s in within if abs(d-nearest)<0.01),key=lambda s:dt(s.get('Start Date'))).get('Locations'))
                eligible=[(d,s) for d,s in within if norm(s.get('Locations'))==nearest_loc]
                chosen_dist,chosen=min(eligible,key=lambda x:dt(x[1].get('Start Date'))); chosen_dist=round(chosen_dist,1)
                explanation=f'Selected earliest eligible physical session at nearest location ({chosen_dist:.1f} miles).'
            elif virtual:
                chosen=min(virtual,key=lambda s:dt(s.get('Start Date'))); explanation=f'No eligible physical session is within {MAX_DISTANCE_MILES:.0f} miles; selected virtual session.' if mapped else 'No eligible mapped physical session; selected virtual session.'
            elif mapped:
                nearest=min(d for d,_ in mapped); rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']=f'Nearest eligible physical session is {nearest:.1f} miles away, over the {MAX_DISTANCE_MILES:.0f}-mile limit, and no virtual session exists.'; continue
            else:
                rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']='Eligible physical sessions exist but their locations cannot be mapped, and no virtual session exists.'; continue
            k=skey(chosen); before=remaining.get(k,0)
            if before<=0: rec['Disposition']='REVIEW_REQUIRED'; rec['Explanation']='Selected session lost its last remaining seat during processing.'; continue
            remaining[k]=before-1
            rec.update({'Disposition':'PROPOSED_SCHEDULE','Selected_Session_WID':id_norm(chosen.get('WID')),'Selected_Reference_ID':clean(chosen.get('Reference ID')),'Selected_Start':chosen.get('Start Date'),'Selected_End':chosen.get('End Date'),'Selected_Location':clean(chosen.get('Locations')) or 'Virtual','Distance_Miles':chosen_dist,'Original_Available_Seats':opening.get(k,0),'Remaining_Seats_After_Reservation':remaining[k],'Workday_Ready':'Yes','Explanation':explanation})
            if prereq_bound: rec['Explanation'] += f" Prerequisite scheduled earlier ({prereq_bound.strftime('%m/%d/%Y')})."
            chosen_st=dt(chosen.get('Start Date')); chosen_en=dt(chosen.get('End Date'))
            if chosen_st and chosen_en and chosen_en > chosen_st:
                busy_by_person[pkey].append({'start':chosen_st,'end':chosen_en,'source':'SELECTED','title':rec['Training_Title'],'event_key':rec['Event_Key']})
            same_run_assignment[(pkey,norm(rec['Training_Title']))]=rec
            reservations.append({'Session_Key':k,'Training_Title':rec['Training_Title'],'Start_Date':chosen.get('Start Date'),'End_Date':chosen.get('End Date'),'Location':clean(chosen.get('Locations')) or 'Virtual','Employee_ID':rec['Employee_ID'],'WID':rec['WID'],'Person_Key':rec.get('Person_Key',pkey),'Event_Key':rec['Event_Key'],'Seats_Before':before,'Seats_After':remaining[k]})

    for r in reqs:
        if r['Disposition']=='PENDING_SCHEDULING': r['Disposition']='REVIEW_REQUIRED'; r['Explanation']='Scheduling engine could not resolve this requirement.'

    req_columns=['Event_Key','Event_Type','Employee_ID','Current_Employee_ID','WID','Person_Key','Worker_Name','Worker_Type','Traveler_Designation','Anchor_Date','Job_Code','Position_Title','Cost_Center_ID','Cost_Center_Title','Sup_Org_ID','Physical_Location','Hiring_Manager_AD','Work_Email','Home_Email','Training_Title','Prerequisite','Topic','Scheduling_Policy','Timing_Modifier','Timing_Min','Timing_Max','Completion_Date','Completion_Training','Existing_Registration_Date','Existing_Session_Start','Existing_Enrollment_Training','Equivalency_Used','Equivalency_Direction','Equivalency_Match_Direction','Selected_Session_WID','Selected_Reference_ID','Selected_Start','Selected_End','Selected_Location','Distance_Miles','Original_Available_Seats','Remaining_Seats_After_Reservation','Prerequisite_Status','Prerequisite_Scheduled_Start','Conflict_Detail','Satisfied_By_Event_Key','Satisfied_By_Session_WID','Disposition','Explanation','Workday_Ready','Override','Override_Reason']
    requirements_df=pd.DataFrame(reqs,columns=req_columns)
    return {
        'events':pd.DataFrame(events),
        'duplicate_events':pd.DataFrame(duplicate_events),
        'existing_workday':pd.DataFrame(existing_workday_rows, columns=['Person_Key','WID','Employee_ID','Training_Title','Registration_Status','Start_Date','End_Date','Location','Course_Offering','Required_For_Current_Role','Scheduling_Impact']),
        'requirements':requirements_df,
        'sessions':pd.DataFrame(sessions),
        'seat_reservations':pd.DataFrame(reservations),
        'summary':dict(Counter(r['Disposition'] for r in reqs)),
        'source_counts':{
            'new_hires':len(new_hires),'job_changes':len(job_changes),
            'staffing_events_raw':raw_event_count,
            'staffing_events_canonical':len(events),
            'duplicate_source_events':len(duplicate_events),
            'history':len(history),'sessions':len(sessions)
        },
        'warning':''
    }

def apply_manual_override(requirements_df, row_index, selected_session, reason):
    df=requirements_df.copy()
    if not reason or not str(reason).strip(): raise ValueError('Override reason is required.')
    for key,val in selected_session.items():
        if key in df.columns: df.at[row_index,key]=val
    df.at[row_index,'Disposition']='PROPOSED_SCHEDULE'; df.at[row_index,'Workday_Ready']='Yes'; df.at[row_index,'Override']='Yes'; df.at[row_index,'Override_Reason']=str(reason).strip(); df.at[row_index,'Explanation']='Manual override: '+str(reason).strip()
    return df
